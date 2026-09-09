"""Lista de trabajo: qué tiene que arreglar una persona, fila a fila.

El hueco que esto cierra
------------------------
Hasta ahora Kalman devolvía tres cosas: un CSV con las filas correctas, otro
con las filas rechazadas y un PDF con recuentos. Eso dice **cuántos** datos
están mal, pero no sirve para arreglarlos.

Un gestor que recibe un CSV con 448 filas rechazadas sigue sin saber qué mirar
en cada una. Tiene que abrirlas y buscar el fallo a ojo, que es exactamente el
trabajo que venía a evitar.

Lo que de verdad necesita es una lista: una línea por problema, con el nombre
del cliente al que hay que llamar, el campo concreto, qué le pasa y qué se
propone. Eso es lo que su equipo abre el lunes por la mañana y va tachando.

Esa lista es el producto. Lo demás es el diagnóstico.
"""

from __future__ import annotations

import io

import pandas as pd

from ..core.dedupe import DuplicateGroup
from ..core.types import CleanReport, Severity

#: Traducción de la gravedad interna a lo que lee una persona.
_GRAVEDAD = {
    Severity.ERROR: "Error",
    Severity.WARNING: "Revisar",
    Severity.INFO: "Corregido",
}

#: Orden de trabajo: primero lo que está demostrablemente mal.
_PRIORIDAD = {"Error": 0, "Revisar": 1, "Corregido": 2}

#: Qué hacer con cada tipo de hallazgo, en imperativo. Sin esto la lista dice
#: cuál es el problema pero no qué se espera de quien la lee.
_ACCION = {
    "tax_id": "Confirmar el identificador fiscal con el cliente.",
    "iban": "Pedir el número de cuenta correcto antes del próximo recibo.",
    "email": "Confirmar la dirección de correo.",
    "phone": "Confirmar el teléfono.",
    "postal_code": "Corregir el código postal.",
    "duplicate": "Decidir qué ficha se conserva y cuál se archiva.",
    "name": "Revisar el nombre del contacto.",
    "company": "Revisar la razón social.",
    "age": "Comprobar el dato con el cliente.",
    "income": "Comprobar el importe con el cliente.",
}

#: Nombre legible de la columna cuando el hallazgo afecta a la fila entera.
_FILA_COMPLETA = "Fila completa"


def _accion_para(rule: str) -> str:
    familia = rule.split(".", 1)[0]
    return _ACCION.get(familia, "Revisar el dato.")


def _identificador(fila: pd.Series, canonical: dict[str, str]) -> tuple[str, str]:
    """Devuelve con qué nombre y con qué documento localizar al cliente.

    Sin esto la lista dice "fila 1.482", que no le sirve a nadie para coger el
    teléfono y llamar.
    """
    nombre = ""
    for campo in ("company", "name"):
        columna = canonical.get(campo)
        if columna and columna in fila.index:
            valor = str(fila[columna]).strip()
            if valor and valor.lower() != "nan":
                nombre = valor
                break

    documento = ""
    columna_doc = canonical.get("tax_id")
    if columna_doc and columna_doc in fila.index:
        valor = str(fila[columna_doc]).strip()
        if valor and valor.lower() != "nan":
            documento = valor

    return nombre, documento


def build_worklist(
    source: pd.DataFrame,
    report: CleanReport,
    duplicates: list[DuplicateGroup] | None = None,
    include_info: bool = False,
) -> pd.DataFrame:
    """Construye la lista de trabajo a partir del informe del motor.

    `source` es el DataFrame tal y como entró, para poder mostrar el valor
    original y el nombre del cliente.

    `include_info` añade lo que Kalman ya ha corregido solo. Va desactivado por
    defecto: una lista de tareas no debe empezar con doscientas líneas de cosas
    que ya están hechas.
    """
    canonical = {v: k for k, v in report.columns_detected.items()}
    duplicados_por_fila: dict[int, int] = {}
    for grupo in duplicates or []:
        for fila in grupo.duplicates:
            duplicados_por_fila[fila] = grupo.survivor

    filas: list[dict[str, object]] = []

    for hallazgo in report.findings:
        gravedad = _GRAVEDAD[hallazgo.severity]
        if gravedad == "Corregido" and not include_info:
            continue

        if 0 <= hallazgo.row < len(source):
            nombre, documento = _identificador(source.iloc[hallazgo.row], canonical)
        else:
            nombre, documento = "", ""

        es_fila_entera = hallazgo.column == "__row__"
        columna = _FILA_COMPLETA if es_fila_entera else hallazgo.column

        if es_fila_entera:
            valor_actual = ""
        elif hallazgo.original is not None:
            valor_actual = str(hallazgo.original)
        elif hallazgo.column in source.columns and 0 <= hallazgo.row < len(source):
            valor_actual = str(source.at[hallazgo.row, hallazgo.column])
        else:
            valor_actual = ""

        filas.append({
            # Se suma 2 porque la fila 0 del DataFrame es la linea 2 del
            # fichero: la 1 son las cabeceras. Asi el numero coincide con lo
            # que ve la persona al abrir el CSV en Excel.
            "Fila en el fichero": hallazgo.row + 2,
            "Cliente": nombre,
            "NIF o CIF": documento,
            "Campo": columna,
            "Valor actual": valor_actual if valor_actual.lower() != "nan" else "",
            "Problema": hallazgo.message,
            "Sugerencia": "" if hallazgo.proposed is None else str(hallazgo.proposed),
            "Qué hacer": _accion_para(hallazgo.rule),
            "Gravedad": gravedad,
            "Confianza": round(hallazgo.confidence, 2),
            "Regla": hallazgo.rule,
        })

    if not filas:
        return pd.DataFrame(
            columns=[
                "Fila en el fichero", "Cliente", "NIF o CIF", "Campo",
                "Valor actual", "Problema", "Sugerencia", "Qué hacer",
                "Gravedad", "Confianza", "Regla",
            ]
        )

    tabla = pd.DataFrame(filas)
    tabla["_orden"] = tabla["Gravedad"].map(_PRIORIDAD)
    tabla = tabla.sort_values(
        ["_orden", "Fila en el fichero", "Campo"], kind="stable"
    ).drop(columns="_orden")
    return tabla.reset_index(drop=True)


#: Ancho de cada columna en el Excel, en caracteres. Una hoja donde hay que
#: ensanchar las columnas a mano antes de poder leerla no parece un producto.
_ANCHOS = {
    "Fila en el fichero": 10,
    "Cliente": 30,
    "NIF o CIF": 13,
    "Campo": 20,
    "Valor actual": 28,
    "Problema": 52,
    "Sugerencia": 28,
    "Qué hacer": 46,
    "Gravedad": 11,
    "Confianza": 10,
    "Regla": 24,
}


def to_excel(worklist: pd.DataFrame, sheet_name: str = "Lista de trabajo") -> bytes:
    """Exporta la lista a Excel, lista para trabajar.

    Una gestoría vive en Excel. Entregarle un CSV crudo la obliga a un paso de
    importación en el que se pierden los ceros iniciales de los códigos
    postales, que es justo uno de los errores que Kalman detecta.

    La hoja sale con la cabecera fijada, filtros puestos y anchos correctos,
    de modo que se pueda ordenar y repartir el trabajo sin tocar nada.
    """
    buffer = io.BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        worklist.to_excel(writer, index=False, sheet_name=sheet_name)
        hoja = writer.sheets[sheet_name]

        from openpyxl.styles import Alignment, Font, PatternFill

        cabecera_fondo = PatternFill("solid", fgColor="0E1726")
        cabecera_letra = Font(color="FFFFFF", bold=True, size=10)

        for celda in hoja[1]:
            celda.fill = cabecera_fondo
            celda.font = cabecera_letra
            celda.alignment = Alignment(vertical="center")

        for indice, nombre in enumerate(worklist.columns, start=1):
            letra = hoja.cell(row=1, column=indice).column_letter
            hoja.column_dimensions[letra].width = _ANCHOS.get(str(nombre), 18)

        # Todo el texto de las celdas se ajusta arriba: los mensajes de
        # problema son largos y centrados verticalmente se leen peor.
        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                celda.alignment = Alignment(vertical="top", wrap_text=False)

        hoja.freeze_panes = "A2"
        if len(worklist) > 0:
            hoja.auto_filter.ref = hoja.dimensions

    return buffer.getvalue()


def summary_line(worklist: pd.DataFrame) -> str:
    """Frase de una línea para encabezar la lista en pantalla."""
    if worklist.empty:
        return "No hay nada que arreglar."

    errores = int((worklist["Gravedad"] == "Error").sum())
    revisar = int((worklist["Gravedad"] == "Revisar").sum())
    clientes = int(worklist["Fila en el fichero"].nunique())

    partes = []
    if errores:
        partes.append(f"{errores} para corregir")
    if revisar:
        partes.append(f"{revisar} para revisar")

    return f"{' y '.join(partes)}, repartidas en {clientes} filas."
